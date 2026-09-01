# -*- coding: utf-8 -*-
"""上传文档的解析与切块模块。

管线：四格式（md/pdf/docx/xlsx）→ 统一 Markdown → 标题切分（面包屑）
→ 递归细切 → 垃圾块过滤 → 父/子块结构（small-to-big）。

PDF 采用混合解析策略：文本页走 pymupdf4llm（免费、毫秒级），
扫描页（文字层稀薄且含图片）渲染位图后走 Vision LLM 兜底
（复用 ai_service 现有 Provider 体系，glm-5.3-flash 等多模态模型）。

块结构约定（与 KnowledgeChunk 模型对应）：
- 短节（<= CHUNK_SIZE）：单块，无父子（parent_id=NULL，有向量）
- 长节：父块存节全文（无向量）+ N 个子块（parent_id 指向父块，有向量）
"""
import base64
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from langchain_text_splitters import (
    MarkdownHeaderTextSplitter,
    RecursiveCharacterTextSplitter,
)

log = logging.getLogger(__name__)

# --- 切块参数（中文按字符估算：800 字符 ≈ 500-600 token，落在检索最佳区间） ---
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
MIN_CHUNK_CHARS = 50  # 短于此的块视为标题孤儿/噪声，丢弃
XLSX_MAX_ROWS_PER_SHEET = 500  # 单 sheet 超出后截断，防止超大表拖垮解析

# --- PDF 混合解析参数（文本层探测 + Vision LLM 兜底） ---
SCAN_PAGE_TEXT_CHARS = 50  # 页面文字层少于此值且含图片 → 视为扫描页，走 Vision 解析

# 支持的上传格式（白名单）
SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx", ".xlsx"}

_HEADER_SPLITTER = MarkdownHeaderTextSplitter(
    headers_to_split_on=[("#", "h1"), ("##", "h2"), ("###", "h3")],
    strip_headers=False,  # 标题保留在正文中，标题词参与 BM25 / embedding
)

_RECURSIVE_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", "。", "；", "，", " ", ""],  # 中文友好
)


@dataclass
class Section:
    """标题切分产出的一个章节。"""

    path: str  # 面包屑，如 "第二章 部署 > 2.1 环境准备"；无标题时为空串
    content: str


@dataclass
class Block:
    """入库前的最终块结构。

    parent_text 为 None 表示单块（整节即一块，无父子结构）。
    """

    path: str
    parent_text: Optional[str] = None
    children: list = field(default_factory=list)


@dataclass
class ChunkingResult:
    filename: str
    markdown: str
    blocks: list = field(default_factory=list)

    @property
    def parent_count(self) -> int:
        return sum(1 for b in self.blocks if b.parent_text)

    @property
    def child_count(self) -> int:
        return sum(len(b.children) for b in self.blocks)


# ---------------------------------------------------------------------------
# 第一层：四格式 → 统一 Markdown
# ---------------------------------------------------------------------------


async def parse_to_markdown(filename: str, raw: bytes) -> str:
    """按扩展名分发到对应解析器，统一产出 Markdown 文本。"""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"不支持的文件格式: {ext}（支持: {', '.join(sorted(SUPPORTED_EXTENSIONS))}）"
        )
    if ext == ".md":
        text = raw.decode("utf-8", errors="replace")
    elif ext == ".pdf":
        text = await _pdf_to_markdown(raw)
    elif ext == ".docx":
        text = _docx_to_markdown(raw)
    else:  # .xlsx
        text = _xlsx_to_markdown(raw)
    if not text.strip():
        raise ValueError("文件解析结果为空（可能是扫描版 PDF 或空文档）")
    return text


async def _vision_page_to_markdown(png_bytes: bytes, page_no: int) -> Optional[str]:
    """单页扫描图 → Markdown（Vision LLM 解析）。

    惰性导入 ai_service 复用现有 Provider 体系；失败返回 None 由调用方标注跳过。
    """
    try:
        from src.chat.services.ai.service import ai_service
        from src.chat.services.ai.providers.base import GenerationConfig

        b64 = base64.b64encode(png_bytes).decode()
        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{b64}"},
                    },
                    {
                        "type": "text",
                        "text": (
                            "这是扫描版文档的一页。请把它完整转录为 Markdown：\n"
                            "1. 按阅读顺序输出正文，标题用 # 层级表示；\n"
                            "2. 表格用 Markdown 管道语法还原；\n"
                            "3. 只转录页面上的真实内容，不要编造或补充任何信息；\n"
                            "4. 只输出 Markdown 本身，不要任何解释。"
                        ),
                    },
                ],
            }
        ]
        # 注意：模型为思考型（reasoning tokens 计入 max_tokens），需给足额度
        result = await ai_service.generate(
            messages,
            config=GenerationConfig(max_output_tokens=4096, temperature=0.1),
        )
        content = (result.content or "").strip()
        if not content:
            log.warning(f"[Vision解析] 第 {page_no} 页返回空内容")
            return None
        log.info(f"[Vision解析] 第 {page_no} 页完成，产出 {len(content)} 字符")
        return content
    except Exception as e:
        log.warning(f"[Vision解析] 第 {page_no} 页解析失败: {e}")
        return None


def _page_is_scan(page) -> bool:
    """探测页面是否为扫描页：文字层稀薄且含图片。"""
    if len(page.get_text().strip()) >= SCAN_PAGE_TEXT_CHARS:
        return False
    return bool(page.get_images(full=True))


async def _pdf_to_markdown(raw: bytes, vision_parser=None) -> str:
    """PDF → Markdown 混合解析。

    逐页探测文字层：正常页用 pymupdf4llm 提取（免费、毫秒级）；
    扫描页（文字稀薄且含图片）渲染为位图后走 Vision LLM 兜底。
    vision_parser 可注入替换（测试用），默认调用 ai_service 视觉解析。
    """
    import pymupdf
    import pymupdf4llm

    if vision_parser is None:
        vision_parser = _vision_page_to_markdown

    doc = pymupdf.open(stream=raw, filetype="pdf")
    try:
        parts: list[str] = []
        for i, page in enumerate(doc):
            if _page_is_scan(page):
                pix = page.get_pixmap(dpi=150)
                png = pix.tobytes("png")
                log.info(
                    f"[PDF混合解析] 第 {i + 1} 页为扫描页（文字层"
                    f" {len(page.get_text().strip())} 字符），转 Vision 解析"
                )
                md = await vision_parser(png, i + 1)
                if md:
                    parts.append(md)
                else:
                    parts.append(f"<!-- 第 {i + 1} 页为扫描页，视觉解析不可用，已跳过 -->")
            else:
                parts.append(pymupdf4llm.to_markdown(doc, pages=[i]))
        return "\n\n".join(parts)
    finally:
        doc.close()


def _docx_to_markdown(raw: bytes) -> str:
    """Word → Markdown：Heading 样式转 # 层级，表格转 MD 表格（python-docx）。"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(raw))
    parts: list[str] = []
    # iter_inner_content 按文档真实顺序遍历段落与表格（python-docx >= 1.1）
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            style_name = (item.style.name or "").lower() if item.style else ""
            text = item.text.strip()
            if not text:
                continue
            if style_name.startswith("heading"):
                try:
                    level = int(style_name.rsplit(" ", 1)[-1])
                except ValueError:
                    level = 2
                level = max(1, min(level, 3))  # ### 以下不再细分，保证切分层级可控
                parts.append(f"{'#' * level} {text}")
            elif style_name.startswith("title"):
                parts.append(f"# {text}")
            else:
                parts.append(text)
        elif isinstance(item, Table):
            parts.append(_table_to_markdown(item))
    return "\n\n".join(parts)


def _table_to_markdown(table) -> str:
    """python-docx / openpyxl 通用：二维行数据 → MD 表格（含表头分隔行）。"""
    rows = []
    for row in table.rows:
        cells = [(_escape_cell(c.text) if hasattr(c, "text") else _escape_cell(str(c))) for c in row.cells]
        rows.append(cells)
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _xlsx_to_markdown(raw: bytes) -> str:
    """Excel → Markdown：每个 sheet 一个 '## Sheet: 名称' 节 + MD 表格（openpyxl）。"""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"## Sheet: {sheet.title}")
            rows = sheet.iter_rows(values_only=True)
            header = None
            count = 0
            lines: list[str] = []
            for row in rows:
                if count >= XLSX_MAX_ROWS_PER_SHEET:
                    lines.append(f"（数据过长，仅保留前 {XLSX_MAX_ROWS_PER_SHEET} 行）")
                    break
                if count == 0:
                    header = [_escape_cell(v) for v in row]
                    lines.append("| " + " | ".join(header) + " |")
                    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                else:
                    lines.append("| " + " | ".join([_escape_cell(v) for v in row]) + " |")
                count += 1
            parts.append("\n".join(lines))
        return "\n\n".join(parts)
    finally:
        workbook.close()


def _escape_cell(value) -> str:
    """单元格 → MD 表格安全文本：竖线转义、换行折叠为空格。"""
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\r\n", " ").replace("\n", " ").strip()


# ---------------------------------------------------------------------------
# 第二层：标题切分 + 面包屑
# ---------------------------------------------------------------------------


def split_sections(md_text: str) -> list[Section]:
    """按 Markdown 标题层级切分，产出带面包屑路径的章节列表。"""
    sections: list[Section] = []
    for doc in _HEADER_SPLITTER.split_text(md_text):
        meta = doc.metadata
        breadcrumb = " > ".join(
            meta.get(key) for key in ("h1", "h2", "h3") if meta.get(key)
        )
        content = doc.page_content.strip()
        if content:
            sections.append(Section(path=breadcrumb, content=content))
    return sections


# ---------------------------------------------------------------------------
# 第三层：递归细切 + 垃圾过滤 → 父/子块
# ---------------------------------------------------------------------------


def _is_table_dominated(text: str) -> bool:
    """节内容以表格为主时禁止细切（表格切碎即失去列关联）。"""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 4:
        return False
    table_lines = sum(1 for ln in lines if ln.lstrip().startswith("|"))
    return table_lines / len(lines) >= 0.6


def build_blocks(sections: list[Section]) -> list[Block]:
    """章节 → 父/子块结构；短节单块，长节拆父子；过滤噪声块。"""
    blocks: list[Block] = []
    for sec in sections:
        # 仅丢弃"带面包屑的短节"（标题孤儿，如只剩一句过渡话）；
        # 无标题的整篇短文（path 为空）是全文，必须保留
        if sec.path and len(sec.content.strip()) < MIN_CHUNK_CHARS:
            continue
        if len(sec.content) <= CHUNK_SIZE or _is_table_dominated(sec.content):
            blocks.append(Block(path=sec.path, parent_text=None, children=[sec.content]))
            continue
        children = [
            c.strip()
            for c in _RECURSIVE_SPLITTER.split_text(sec.content)
            if len(c.strip()) >= MIN_CHUNK_CHARS
        ]
        if not children:
            continue
        blocks.append(Block(path=sec.path, parent_text=sec.content, children=children))
    return blocks


def chunk_markdown(md_text: str, source_name: str = "document") -> ChunkingResult:
    """对已解析好的 Markdown/纯文本切块（入库流程入口）。

    纯文本（无 # 标题）会整体成为一个无面包屑的 section，
    短文本走单块路径，行为与不分块时代一致。
    """
    sections = split_sections(md_text)
    blocks = build_blocks(sections)
    if not blocks:
        raise ValueError("切块结果为空：文档可能只有标题或噪声内容")
    result = ChunkingResult(filename=source_name, markdown=md_text, blocks=blocks)
    log.info(
        f"[切块] {source_name}: {len(sections)} 个章节 → "
        f"{len(blocks)} 块（父块 {result.parent_count} / 子块 {result.child_count}）"
    )
    return result


async def chunk_document(filename: str, raw: bytes) -> ChunkingResult:
    """上传文件字节流的完整管线入口：解析 → 切块。"""
    markdown = await parse_to_markdown(filename, raw)
    return chunk_markdown(markdown, source_name=filename)
